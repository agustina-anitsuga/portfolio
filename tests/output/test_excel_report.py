# -*- coding: utf-8 -*-
import openpyxl

from doubles import USD_MARKET, make_holding, make_instrument
from portfolio_dashboard.market import Market
from portfolio_dashboard.marketdata.fx_rate import FxRate
from portfolio_dashboard.output.excel_report import ExcelReport
from portfolio_dashboard.portfolio.market_report import MarketReport
from portfolio_dashboard.portfolio.snapshot import PortfolioSnapshot


def snapshot(with_positions=True):
    reports = {key: MarketReport(Market.get(key), []) for key in Market.keys()}
    if with_positions:
        reports["usd"] = MarketReport(USD_MARKET, [make_holding(instrument=make_instrument("AAA"))])
    return PortfolioSnapshot(reports=reports, fx=FxRate(1000.0, "test"))


def saved(tmp_path, snap=None):
    path = tmp_path / "out.xlsx"
    ExcelReport(snap or snapshot()).save(path)
    return openpyxl.load_workbook(path)


def test_there_is_a_dashboard_and_a_detail_sheet_per_market(tmp_path):
    names = saved(tmp_path).sheetnames
    for key in Market.keys():
        assert f"dashboard-{key}" in names
        assert f"portfolio-{key}" in names


def test_the_general_sheet_comes_first(tmp_path):
    """It is the summary, so it should be the tab that opens."""
    assert saved(tmp_path).sheetnames[0] == "dashboard-general"


def test_the_markets_keep_the_dashboard_order(tmp_path):
    names = [name for name in saved(tmp_path).sheetnames if name.startswith("portfolio-")]
    assert names == [f"portfolio-{key}" for key in Market.keys()]


def test_the_default_empty_sheet_is_removed(tmp_path):
    assert "Sheet" not in saved(tmp_path).sheetnames


def test_the_positions_reach_the_detail_sheet(tmp_path):
    sheet = saved(tmp_path)["portfolio-usd"]
    assert sheet.cell(row=2, column=1).value == "AAA"


def test_an_empty_portfolio_still_saves(tmp_path):
    """Nothing to report is not a reason to crash."""
    workbook = saved(tmp_path, snapshot(with_positions=False))
    assert workbook["dashboard-general"]["A1"].value.startswith("Dashboard General")

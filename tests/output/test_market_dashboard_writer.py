# -*- coding: utf-8 -*-
import openpyxl
import pytest

from doubles import USD_MARKET, make_holding, make_instrument, make_position, make_quote
from portfolio.output.market_dashboard_writer import MarketDashboardWriter
from portfolio.portfolio.market_report import MarketReport


def holding(key="AAA", cost_ars=1000.0, price_ars=200.0):
    return make_holding(instrument=make_instrument(key),
                        position=make_position(units=10.0, cost_ars=cost_ars, cost_usd=10.0),
                        quote=make_quote(price_ars=price_ars, price_usd=price_ars / 100.0))


def write(holdings):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    report = MarketReport(USD_MARKET, list(holdings))
    return MarketDashboardWriter(workbook).write("dashboard-usd", report)


def test_the_title_names_the_market():
    assert write([holding()])["A1"].value == "Dashboard - USD"


def test_the_ars_kpis_go_in_the_first_two_columns():
    sheet = write([holding(cost_ars=1000.0, price_ars=200.0)])
    assert sheet["A3"].value == "Invertido (ARS)"
    assert sheet["B3"].value == pytest.approx(1000.0)
    assert sheet["A4"].value == "Valor Actual (ARS)"
    assert sheet["B4"].value == pytest.approx(2000.0)


def test_the_usd_kpis_go_next_to_them():
    sheet = write([holding()])
    assert sheet["C3"].value == "Invertido (USD)"
    assert sheet["D3"].value == pytest.approx(10.0)


def test_the_percentage_kpi_is_stored_as_a_fraction():
    sheet = write([holding(cost_ars=1000.0, price_ars=200.0)])
    assert sheet["A6"].value == "P&L No Realizado % (ARS)"
    assert sheet["B6"].value == pytest.approx(1.0)      # +100%
    assert sheet["B6"].number_format == "0.00%"


def test_the_top_positions_table_lists_ticker_and_amounts():
    sheet = write([holding("AAA"), holding("BBB")])
    assert [sheet.cell(row=3, column=column).value for column in (6, 7, 8)] == \
        ["Ticker", "Invertido (ARS)", "Valor Actual (ARS)"]
    assert sheet.cell(row=4, column=6).value in ("AAA", "BBB")


def test_the_table_is_capped_at_twelve_positions():
    sheet = write([holding(f"T{index:02d}") for index in range(20)])
    tickers = [sheet.cell(row=row, column=6).value for row in range(4, 30)]
    assert len([t for t in tickers if t]) == 12


def test_an_unpriced_position_contributes_zero_to_the_chart_table():
    from portfolio.marketdata.quote import Quote
    unpriced = make_holding(instrument=make_instrument("ZZZ"), quote=Quote.unavailable("sin datos"))
    sheet = write([unpriced])
    assert sheet.cell(row=4, column=8).value == 0


def test_a_bar_chart_is_attached():
    sheet = write([holding()])
    assert len(sheet._charts) == 1
    assert sheet._charts[0].title is not None


def test_an_empty_market_still_produces_a_readable_sheet():
    sheet = write([])
    assert sheet["A1"].value == "Dashboard - USD"
    assert sheet["B3"].value is None      # nothing invested, nothing reported

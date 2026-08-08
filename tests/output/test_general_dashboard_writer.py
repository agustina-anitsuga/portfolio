# -*- coding: utf-8 -*-
import openpyxl
import pytest

from doubles import CEDEARS_MARKET, USD_MARKET, make_holding, make_instrument, make_quote
from portfolio_dashboard.output.general_dashboard_writer import HEADERS, GeneralDashboardWriter
from portfolio_dashboard.portfolio.market_report import MarketReport


def write(reports):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    return GeneralDashboardWriter(workbook).write(reports)


def report(market, *holdings):
    return MarketReport(market, list(holdings))


def holding(key, price_ars, market=USD_MARKET):
    return make_holding(market=market, instrument=make_instrument(key),
                        quote=make_quote(price_ars=price_ars, price_usd=price_ars / 1000.0))


def test_the_sheet_is_titled_and_has_the_headers():
    sheet = write({"usd": report(USD_MARKET, holding("AAA", 2000.0))})
    assert sheet["A1"].value == "Dashboard General - Resultado por Producto"
    assert [cell.value for cell in sheet[3]] == HEADERS


def test_every_market_contributes_its_rows():
    sheet = write({
        "usd": report(USD_MARKET, holding("AAA", 2000.0)),
        "cedears": report(CEDEARS_MARKET, holding("BBB", 3000.0, CEDEARS_MARKET)),
    })
    tickers = {sheet.cell(row=row, column=1).value for row in (4, 5)}
    assert tickers == {"AAA", "BBB"}


def test_each_row_names_the_portfolio_it_belongs_to():
    sheet = write({"cedears": report(CEDEARS_MARKET, holding("BBB", 3000.0, CEDEARS_MARKET))})
    assert sheet.cell(row=4, column=2).value == "Cedears"


def test_rows_are_sorted_by_result_with_the_best_first():
    sheet = write({"usd": report(USD_MARKET, holding("LOSS", 100.0), holding("WIN", 9000.0))})
    assert [sheet.cell(row=row, column=1).value for row in (4, 5)] == ["WIN", "LOSS"]


def test_percentages_are_stored_as_fractions():
    sheet = write({"usd": report(USD_MARKET, holding("AAA", 2000.0))})
    assert sheet.cell(row=4, column=6).value == pytest.approx(1.0)   # +100%
    assert sheet.cell(row=4, column=6).number_format == "0.00%"


def test_the_result_columns_get_the_colour_scale():
    sheet = write({"usd": report(USD_MARKET, holding("AAA", 2000.0))})
    ranges = {str(formatting.sqref) for formatting in sheet.conditional_formatting}
    assert ranges == {"E4", "F4", "I4", "J4"}


def test_the_header_row_is_frozen_and_filterable():
    sheet = write({"usd": report(USD_MARKET, holding("AAA", 2000.0))})
    assert sheet.freeze_panes == "A4"
    assert sheet.auto_filter.ref == "A3:J4"


def test_a_portfolio_with_no_positions_adds_nothing():
    sheet = write({"usd": report(USD_MARKET)})
    assert sheet.cell(row=4, column=1).value is None

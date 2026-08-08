# -*- coding: utf-8 -*-
import openpyxl
import pytest

from portfolio_dashboard.output.portfolio_sheet_writer import COLUMNS, PortfolioSheetWriter

ROW = {
    "key": "AAA", "name": "Alpha Corp", "sector": "Technology", "units": 10.0,
    "units_sold": 4.0, "price_source": "PPI (en vivo)",
    "avg_cost_ars": 100.0, "price_ars": 150.0, "invested_ars": 1000.0, "value_ars": 1500.0,
    "pl_abs_ars": 500.0, "pl_pct_ars": 50.0, "cost_of_sales_ars": 400.0,
    "income_from_sales_ars": 600.0, "realized_abs_ars": 200.0, "realized_pct_ars": 50.0,
    "avg_cost_usd": 1.0, "price_usd": 1.5, "invested_usd": 10.0, "value_usd": 15.0,
    "pl_abs_usd": 5.0, "pl_pct_usd": 50.0, "cost_of_sales_usd": 4.0,
    "income_from_sales_usd": 6.0, "realized_abs_usd": 2.0, "realized_pct_usd": 50.0,
}


def write(rows):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    sheet = PortfolioSheetWriter(workbook).write("portfolio-usd", rows)
    return sheet


def test_the_first_row_holds_the_headers():
    sheet = write([ROW])
    assert [cell.value for cell in sheet[1]] == [column.header for column in COLUMNS]


def test_each_value_lands_in_its_column():
    sheet = write([ROW])
    values = {column.key: sheet.cell(row=2, column=index).value
              for index, column in enumerate(COLUMNS, start=1)}
    assert values["key"] == "AAA"
    assert values["invested_ars"] == 1000.0
    assert values["value_usd"] == 15.0


def test_percentages_are_stored_as_fractions_with_a_percent_format():
    sheet = write([ROW])
    index = [column.key for column in COLUMNS].index("pl_pct_ars") + 1
    cell = sheet.cell(row=2, column=index)
    assert cell.value == pytest.approx(0.5)
    assert cell.number_format == "0.00%"


def test_money_columns_carry_a_thousands_format():
    sheet = write([ROW])
    index = [column.key for column in COLUMNS].index("invested_ars") + 1
    assert sheet.cell(row=2, column=index).number_format == "#,##0.00"


def test_missing_values_leave_the_cell_empty():
    sheet = write([{"key": "AAA"}])
    index = [column.key for column in COLUMNS].index("value_ars") + 1
    assert sheet.cell(row=2, column=index).value is None


def test_the_header_row_is_frozen_and_filterable():
    sheet = write([ROW])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:Z2"


def test_the_pl_columns_get_a_colour_scale():
    sheet = write([ROW])
    ranges = {str(formatting.sqref) for formatting in sheet.conditional_formatting}
    # the four P&L columns of each currency, and only those
    assert ranges == {"K2", "L2", "O2", "P2", "U2", "V2", "Y2", "Z2"}


def test_an_empty_sheet_still_has_headers_and_a_valid_filter():
    sheet = write([])
    assert sheet.cell(row=1, column=1).value == "Ticker"
    assert sheet.auto_filter.ref == "A1:Z1"


def test_every_column_gets_its_width():
    sheet = write([ROW])
    assert sheet.column_dimensions["A"].width == COLUMNS[0].width
